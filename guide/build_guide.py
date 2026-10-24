# -*- coding: utf-8 -*-
"""Сборка простого руководства: картинки, схемы, HTML и Word."""
import html
import os
import shutil
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = r"C:\Users\Sergei\AppData\Local\Temp\cursor\screenshots"
OUT = os.path.join(ROOT, "guide", "images")
os.makedirs(OUT, exist_ok=True)

FONT = r"C:\Windows\Fonts\segoeui.ttf"
FONTB = r"C:\Windows\Fonts\segoeuib.ttf"

PULT = "https://script.google.com/macros/s/AKfycbyrI9lD_y0La2LAlkWhrR0cra0V4yZSG6zeRDK72NFLlQYRrtsirsUhARjVqVBQTDDscw/exec"
BOOK = "https://docs.google.com/spreadsheets/d/1sjnqaU4-uJD_kjFwZWjnofPd25vQjtCqstktGWhXbe8/edit"
GID = {
    "Памятка": 311012353,
    "Образец": 959991302,
    "Витрина": 30284276,
    "Графики": 390188976,
    "Правки": 1074985368,
    "Журнал": 89825155,
}
GOODS = "https://docs.google.com/spreadsheets/d/1mMZlLOu9XP2BTzs_frWq6uAKzoqC_TC3bBJ1JmWzlyY/edit"
CASH = "https://docs.google.com/spreadsheets/d/1SlIUkncDyrhpvs7FPmtWY6zvX684nbiO-s0bVrzFOGQ/edit"
MS = "https://online.moysklad.ru/"


def sheet(name):
    return BOOK + "?gid=%s#gid=%s" % (GID[name], GID[name])


def font(size, bold=False):
    return ImageFont.truetype(FONTB if bold else FONT, size)


def wrap(draw, text, f, width):
    lines, cur = [], ""
    for word in text.split():
        trial = word if not cur else cur + " " + word
        if draw.textlength(trial, font=f) <= width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines


def rounded(draw, box, fill, outline=None, radius=18, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def center_text(draw, text, f, box, fill, gap=6):
    x, y, x2, y2 = box
    lines = wrap(draw, text, f, (x2 - x) - 24)
    lh = f.size + gap
    total = lh * len(lines) - gap
    yy = y + ((y2 - y) - total) / 2
    for line in lines:
        tw = draw.textlength(line, font=f)
        draw.text((x + ((x2 - x) - tw) / 2, yy), line, font=f, fill=fill)
        yy += lh


def arrow(draw, x1, y1, x2, y2, color="#202124"):
    draw.line((x1, y1, x2, y2), fill=color, width=5)
    draw.polygon([(x2, y2), (x2 - 14, y2 - 8), (x2 - 14, y2 + 8)], fill=color)


def save_diagrams():
    ink = "#202124"
    green = "#188038"
    green_bg = "#E6F4EA"
    amber = "#E37400"
    amber_bg = "#FEF7E0"
    blue = "#1A73E8"
    blue_bg = "#E8F0FE"
    gray = "#5F6368"
    gray_bg = "#F1F3F4"
    white = "#FFFFFF"

    # 1. Три двери
    im = Image.new("RGB", (1400, 420), white)
    d = ImageDraw.Draw(im)
    cards = [
        (40, "1", "Пульт", "Большие цифры.\nСюда смотреть сначала.", green, green_bg),
        (490, "2", "Таблица", "Все листы.\nЗдесь лежат строки.", blue, blue_bg),
        (940, "3", "Графики", "Те же цифры\nстолбиками.", amber, amber_bg),
    ]
    title = font(28, True)
    d.text((40, 24), "Три места. Больше для обычного дня не нужно.", font=title, fill=ink)
    for x, num, name, text, color, bg in cards:
        rounded(d, (x, 90, x + 420, 380), bg, color, 22, 4)
        d.ellipse((x + 24, 114, x + 78, 168), fill=color)
        nf = font(28, True)
        d.text((x + 40, 122), num, font=nf, fill=white)
        d.text((x + 96, 122), name, font=font(32, True), fill=ink)
        yy = 200
        for line in text.split("\n"):
            d.text((x + 28, yy), line, font=font(24), fill=ink)
            yy += 40
    im.save(os.path.join(OUT, "doors.png"))

    # 2. Поток
    im = Image.new("RGB", (1400, 520), white)
    d = ImageDraw.Draw(im)
    d.text((40, 20), "Откуда берётся цифра на пульте", font=font(28, True), fill=ink)
    boxes = [
        (40, 160, 300, 300, "Таблицы\nсотрудников", blue_bg, blue),
        (380, 80, 640, 220, "МойСклад", gray_bg, gray),
        (560, 280, 860, 440, "Робот\nкаждое утро", amber_bg, amber),
        (960, 160, 1360, 340, "Витрина,\nпотом пульт\nи графики", green_bg, green),
    ]
    for x, y, x2, y2, text, bg, color in boxes:
        rounded(d, (x, y, x2, y2), bg, color, 20, 4)
        center_text(d, text.replace("\n", " "), font(26, True), (x, y, x2, y2), ink)
    arrow(d, 300, 230, 560, 340, blue)
    arrow(d, 640, 180, 700, 280, gray)
    arrow(d, 860, 360, 960, 280, green)
    d.text((40, 470), "Робот сам переписывает свод. Руками в «Витрине» цифру не держат.", font=font(22), fill=gray)
    im.save(os.path.join(OUT, "flow.png"))

    # 3. Прибыль
    im = Image.new("RGB", (1400, 420), white)
    d = ImageDraw.Draw(im)
    d.text((40, 20), "Как получается прибыль", font=font(28, True), fill=ink)
    items = [
        (40, "Поступления", "15 000", green_bg, green),
        (360, "минус реклама", "0", gray_bg, gray),
        (680, "Прибыль", "15 000", green_bg, green),
        (1020, "Списание рядом", "2 600", amber_bg, amber),
    ]
    for x, name, num, bg, color in items:
        rounded(d, (x, 100, x + 320, 300), bg, color, 20, 4)
        d.text((x + 24, 130), name, font=font(22, True), fill=ink)
        d.text((x + 24, 190), num, font=font(48, True), fill=color)
    d.text((40, 340), "Прибыль = поступления минус реклама. Списание на карточке есть, из прибыли его не вычитают.", font=font(22), fill=ink)
    d.text((40, 378), "Эти числа учебные. На них написано ТЕСТ. Это не деньги в кассе.", font=font(22), fill=amber)
    im.save(os.path.join(OUT, "profit.png"))

    # 4. Цвета
    im = Image.new("RGB", (1400, 460), white)
    d = ImageDraw.Draw(im)
    d.text((40, 20), "Цвет вкладки внизу таблицы", font=font(28, True), fill=ink)
    rows = [
        (90, green, green_bg, "Зелёный", "Каждый день", "Памятка, франшиза, источники, витрина, графики, правки"),
        (175, amber, amber_bg, "Жёлтый", "Когда цифра странная", "Журнал и бланки"),
        (260, blue, blue_bg, "Синий", "Только для проверки следа", "Сырьё и вход 1С"),
        (345, gray, gray_bg, "Серый", "В обычный день не открывать", "Словарь полей и служебное"),
    ]
    for y, color, bg, name, who, what in rows:
        rounded(d, (40, y, 1360, y + 72), bg, color, 14, 3)
        d.rounded_rectangle((58, y + 18, 150, y + 54), radius=10, fill=color)
        d.text((170, y + 18), name, font=font(24, True), fill=ink)
        d.text((430, y + 18), who, font=font(22, True), fill=ink)
        d.text((820, y + 20), what, font=font(20), fill=ink)
    im.save(os.path.join(OUT, "colors.png"))

    # 5. Правка
    im = Image.new("RGB", (1400, 360), white)
    d = ImageDraw.Draw(im)
    d.text((40, 16), "Как заменить одну сумму и не потерять её", font=font(28, True), fill=ink)
    steps = [
        "1. На пульте открыть источник",
        "2. Вкладка «Корректировка»",
        "3. Дважды щёлкнуть сумму и вписать новую",
        "4. Робот запомнит её на листе «Правки»",
    ]
    for i, text in enumerate(steps):
        x = 30 + i * 345
        rounded(d, (x, 90, x + 320, 280), green_bg if i == 3 else blue_bg, green if i == 3 else blue, 18, 3)
        center_text(d, text, font(24, True), (x + 8, 100, x + 312, 270), ink)
        if i < 3:
            arrow(d, x + 322, 185, x + 342, 185)
    d.text((40, 310), "Чужая таблица сотрудника при этом остаётся как была.", font=font(22), fill=gray)
    im.save(os.path.join(OUT, "edit.png"))

    # 6. День
    im = Image.new("RGB", (1400, 340), white)
    d = ImageDraw.Draw(im)
    d.text((40, 16), "Обычное утро", font=font(28, True), fill=ink)
    day = [
        ("1", ["Открыть", "пульт"]),
        ("2", ["Прочитать", "пять карточек"]),
        ("3", ["Заглянуть", "в графики"]),
        ("4", ["Если странно —", "открыть журнал"]),
    ]
    for i, (num, lines) in enumerate(day):
        x = 30 + i * 345
        rounded(d, (x, 80, x + 325, 300), green_bg, green, 18, 3)
        d.ellipse((x + 126, 104, x + 186, 164), fill=green)
        nw = d.textlength(num, font=font(26, True))
        d.text((x + 156 - nw / 2, 114), num, font=font(26, True), fill=white)
        yy = 186
        for line in lines:
            tw = d.textlength(line, font=font(24, True))
            d.text((x + (325 - tw) / 2, yy), line, font=font(24, True), fill=ink)
            yy += 36
    im.save(os.path.join(OUT, "day.png"))


def crop_shots():
    jobs = {
        "console": ("guide-console.png", 0.50),
        "charts": ("guide-charts-top.png", 0.96),
        "charts-bar": ("guide-charts.png", 0.46),
        "memo": ("guide-memo.png", 0.96),
        "sample": ("guide-sample.png", 0.38),
        "vitrina": ("guide-vitrina-test.png", 0.42),
        "money": ("guide-vitrina-money2.png", 0.46),
        "edit-sheet": ("guide-corrections2.png", 0.40),
        "journal": ("guide-journal-rows.png", 0.58),
    }
    for name, (src, ratio) in jobs.items():
        im = Image.open(os.path.join(SRC, src)).convert("RGB")
        h = max(1, int(im.height * ratio))
        im.crop((0, 0, im.width, h)).save(os.path.join(OUT, name + ".png"), quality=90)


def blocks():
    return [
        ("h1", "Как пользоваться пультом РемМонтаж"),
        ("lead", "Это инструкция простыми словами. Здесь написано, куда нажимать, что означают цифры и какие листы можно не открывать."),
        ("h2", "1. С чего начать"),
        ("p", "Откройте браузер Chrome. Вверху справа должна быть почта alexsklad083@gmail.com. С другой почты пульт не откроется."),
        ("p", "Дальше нужны только три ссылки. Их можно сохранить в закладки."),
        ("img", "doors.png", "Три входа: пульт, таблица и графики"),
        ("links", [
            ("Открыть пульт", PULT),
            ("Открыть таблицу «Мастер-реестр РемМонтаж»", BOOK),
            ("Открыть лист «Графики»", sheet("Графики")),
        ]),
        ("p", "Сверху на пульте может быть голубая полоска: приложение создано пользователем Google. Это обычное предупреждение Google. Крестик справа её закрывает."),
        ("h2", "2. Откуда берётся цифра"),
        ("img", "flow.png", "Путь строки: таблицы и МойСклад, робот, витрина, пульт"),
        ("p", "Сотрудники ведут свои таблицы. Складская программа МойСклад хранит карточки фирм и людей. Робот раз в день, около 07:00 по Москве, читает включённые листы и складывает строки в лист «Витрина». Пульт и графики показывают уже эту витрину."),
        ("p", "Сейчас фирма в реестре одна: ООО РемМонтаж, код FR-001. Статус франшизы: ОТКЛЮЧЕН. Пока статус такой, утренний робот делает копию книги и новые строки не забирает. Учебный образец на пульте при этом виден."),
        ("h2", "3. Что видно на пульте"),
        ("img", "console.png", "Пульт: пять карточек, список слева пустой"),
        ("p", "Слева пусто. Так и задумано, пока франшиза выключена. Справа пять карточек. Сейчас на них учебные числа:"),
        ("ul", [
            "Строк в своде — 2. Это две учебные оплаты с пометкой ТЕСТ.",
            "Поступления — 15 000,00. Деньги, которые пришли.",
            "Списания — 2 600,00. Деньги, которые ушли.",
            "Расход Google Ads — 0,00. Рекламный кабинет есть, кампании и карты нет, выгрузка расхода не сохранена.",
            "Прибыль до себестоимости — 15 000,00.",
        ]),
        ("img", "profit.png", "Прибыль равна поступлениям минус реклама. Списание стоит отдельно."),
        ("p", "Себестоимость товаров из прибыли не вычитается. Поэтому карточка так и называется: прибыль до себестоимости."),
        ("p", "Под карточками написано про рекламный кабинет 868-944-6535. Валюта там доллар, пояс Минск. Менять это самими руками не нужно."),
        ("h2", "4. Что делать утром"),
        ("img", "day.png", "Четыре утренних шага"),
        ("ol", [
            "Откройте пульт по ссылке из первого раздела.",
            "Прочитайте пять карточек. Запомните поступления, списание и прибыль.",
            "Откройте лист «Графики». Столбики должны совпасть с карточками.",
            "Если число выглядит странно, откройте лист «Журнал» и читайте нижние строки.",
        ]),
        ("p", "Пока франшиза выключена, слева на пульте имён нет. Это не поломка. Большие цифры сверху всё равно живые: в них входит учебный образец."),
        ("h2", "5. Цвета вкладок"),
        ("p", "Внизу таблицы у каждого листа свой цвет. Цвет подсказывает, открывать лист сегодня или нет."),
        ("img", "colors.png", "Зелёный каждый день, жёлтый при странной цифре, синий для следа, серый не для смены"),
        ("p", "Порядок листов слева направо: Памятка, Образец, Управление Франшизами, Источники, Витрина, Графики, Правки, Журнал, Бланки, Сырьё, Вход 1С, Словарь полей, Служебное."),
        ("h2", "6. Лист «Памятка»"),
        ("p", "Это короткая шпаргалка внутри самой таблицы. Зелёные строки — ежедневные листы. Жёлтые — когда сбор выглядит странно. Синие — технический след. Серые в смену не открывают."),
        ("img", "memo.png", "Лист Памятка"),
        ("links", [("Открыть «Памятку»", sheet("Памятка"))]),
        ("h2", "7. Лист «Образец»"),
        ("p", "Две жёлтые строки. На них написано ТЕСТ. Первая: 1 октября 2026, поступление 15 000. Вторая: 2 октября 2026, списание 2 500. Это пример для глаза. В кассе таких денег нет."),
        ("img", "sample.png", "Две учебные строки на листе Образец"),
        ("links", [("Открыть «Образец»", sheet("Образец"))]),
        ("h2", "8. Лист «Витрина»"),
        ("p", "Сюда робот складывает все строки. Смотреть нужно жёлтые столбцы: поступление и списание. Внизу листа лежат учебные строки. У оплаты заказчика в поступлении 15 000. У оплаты поставщику в списании 2 600. Пять строк «ИМИТАЦИЯ» денег не несут: это пример товаров, не склад."),
        ("img", "vitrina.png", "Внизу витрины учебные имена"),
        ("img", "money.png", "Жёлтые столбцы: поступление 15 000 и списание 2 600"),
        ("p", "Выше этих строк много настоящих имён контрагентов и сотрудников. Денег в тех строках нет. Пульт их в сумму 15 000 не кладёт, потому что франшиза выключена, а образец в сумму входит всегда."),
        ("p", "Цифру в витрине насовсем не правят. Следующий сбор запишет лист заново. Постоянная замена живёт на листе «Правки»."),
        ("links", [("Открыть «Витрину»", sheet("Витрина"))]),
        ("h2", "9. Лист «Графики»"),
        ("p", "Те же пять чисел, что на пульте, плюс два рисунка. Левый столбик и правый столбик высокие: это 15 000. Средний низкий: это списание 2 600. Нижний рисунок показывает две учебные строки по именам. Синий цвет — поступление, красный — списание."),
        ("img", "charts.png", "Лист Графики: числа и столбики. Синий — поступление, красный — списание."),
        ("p", "Жёлтая записка на листе говорит то же самое: это не факт оплаты. Когда в витрине появятся новые деньги включённой франшизы, графики пересчитаются сами."),
        ("links", [("Открыть «Графики»", sheet("Графики"))]),
        ("h2", "10. Лист «Правки»"),
        ("p", "Здесь одна учебная замена. В источнике списание поставщику равно 2 500. На витрине и на пульте оно равно 2 600. Разницу держит этот лист. Робот при каждом сборе снова подставляет 2 600."),
        ("img", "edit-sheet.png", "Лист Правки: списание 2 600 от 23 октября 2026"),
        ("img", "edit.png", "Четыре шага, чтобы сохранить свою сумму"),
        ("p", "Свой двойной щелчок на пульте получится тогда, когда слева появится имя источника. Пока франшиза выключена, список пустой, и новую правку с пульта не из чего выбрать. Уже записанная строка на листе «Правки» при этом работает."),
        ("p", "Таблица сотрудника от правки не меняется. Меняется только число в своде."),
        ("links", [("Открыть «Правки»", sheet("Правки"))]),
        ("h2", "11. Лист «Журнал»"),
        ("p", "Это дневник робота. Новые строки снизу. Цвет строки важнее длинного текста."),
        ("ul", [
            "Зелёная, уровень «ок» — шаг записан.",
            "Жёлтая, уровень «предупреждение» — сбор жив, кусок пропущен. Часто это незнакомый столбец. Его пропустили, остальное записали.",
            "Синяя, уровень «заметка» — это не поломка. Робот объясняет, почему шаг тихий.",
            "Красная, уровень «ошибка» — шаг не записан. Такую строку не стирают, пока причина не закрыта.",
        ]),
        ("img", "journal.png", "Журнал: зелёные и жёлтые строки утра 22 октября"),
        ("p", "На снимке нижняя синяя строка от 22 октября говорит, что проект Google Cloud не задан. Это старая запись. Сейчас проект задан, а живые цифры смотрят на пульте и на листе «Графики»."),
        ("p", "Строка «0 строк. Файл карточек ещё не передан» значит: настоящего файла товаров ещё нет. Пять жёлтых строк «ИМИТАЦИЯ» — только пример."),
        ("links", [("Открыть «Журнал»", sheet("Журнал"))]),
        ("h2", "12. Остальные листы короткими словами"),
        ("ul", [
            "Управление Франшизами. Одна рабочая строка: FR-001, ООО РемМонтаж, статус ОТКЛЮЧЕН. Короткое имя организации МойСклад: ooorusuglneftgaz. Длинный номер рядом — не пароль и не ключ. Столбец с ключом не стирать и никому в письме не пересылать.",
            "Источники. Список таблиц, которые робот умеет читать. Статус «Включать» значит «эта таблица нужна». Пока франшиза выключена, утро всё равно их пропускает. Подпись FR-002 сбор не запускает: отдельной рабочей франшизы с таким кодом нет.",
            "Бланки. Проверка, есть ли в чужой таблице нужный столбец. Сюда не кладут сканы и договоры.",
            "Сырьё. Черновик до витрины. Открывать, только если цифра спорит с источником. Руками не править.",
            "Вход 1С. Архив посылок для робота. Сами в витрину они не попадают.",
            "Словарь полей. Перевод заголовков. В смену не открывать и имена полей не переименовывать.",
            "Служебное. Настройки проекта. В смену не открывать. Регион EU не менять. Оплату Google не включать.",
        ]),
        ("h2", "13. Если картина странная"),
        ("ul", [
            "На пульте ноль строк, хотя образец на месте. Обновите страницу. Образец должен давать 2 строки, 15 000 и 2 600.",
            "На графике другие числа, чем на пульте. Подождите минуту: таблица пересчитывает формулы. Потом обновите оба окна.",
            "В журнале жёлтая строка про незнакомый столбец. Сбор при этом жив. Лишние столбцы робот пропускает.",
            "В журнале красная строка. Её не удаляют. Смотрят текст и лист, который в ней назван.",
            "Статус франшизы стал «ОШИБКА ТОКЕНА». Сбор склада остановился. Ключ в таблице не переписывают от руки.",
            "Слева на пульте пусто. При статусе ОТКЛЮЧЕН так и должно быть.",
        ]),
        ("h2", "14. В обычный день эти места не трогать"),
        ("ul", [
            "Ключ в листе «Управление Франшизами».",
            "Серые листы «Словарь полей» и «Служебное».",
            "Кнопка смены тарифа в МойСклад и любая оплата Google.",
            "Кнопка новой рекламной кампании.",
            "Чужие таблицы сотрудников: робот их читает и сам туда новые суммы не пишет.",
            "Жёлтые строки ТЕСТ и ИМИТАЦИЯ. Их не принимают за кассу и за склад. Перед настоящим файлом товаров пять жёлтых строк «ИМИТАЦИЯ» сначала убирают.",
        ]),
        ("h2", "15. Все ссылки"),
        ("links", [
            ("Пульт", PULT),
            ("Мастер-реестр", BOOK),
            ("Памятка", sheet("Памятка")),
            ("Образец", sheet("Образец")),
            ("Витрина", sheet("Витрина")),
            ("Графики", sheet("Графики")),
            ("Правки", sheet("Правки")),
            ("Журнал", sheet("Журнал")),
            ("Товары, пять жёлтых строк-примеров", GOODS),
            ("Движение денег, месяц пока пустой", CASH),
            ("МойСклад", MS),
        ]),
        ("p", "В МойСклад заходят под своей учётной записью фирмы. Организация там называется «РУСУГЛЕНЕФТЕГАЗ О.». На экране кассы и франшизы имя остаётся «ООО РемМонтаж». Пробный срок МойСклад заканчивается 2 ноября 2026 в 12:32. Продлевать его или нет, решает заказчик. Кнопку выбора тарифа сами не нажимают."),
        ("h2", "16. Что на экране учебное, а что ещё впереди"),
        ("p", "Уже можно открывать пульт, реестр и графики. Ночная копия книги пишется сама, в папку Диска «Бэкапы РемМонтаж». Хранятся 14 последних копий."),
        ("p", "Числа 15 000 и 2 600 — учебные. Настоящий файл товаров ещё не положен. Настоящих платежей в МойСклад нет. Расход рекламы равен нулю. Отдельный красивый отчёт Google из этой страны не открывается: его заменяет лист «Графики» и пульт."),
    ]


def render_html(parts):
    css = """
    body { margin: 0; background: #f6f8f6; color: #202124; font: 20px/1.45 "Segoe UI", Arial, sans-serif; }
    main { max-width: 980px; margin: 0 auto; padding: 28px 20px 80px; }
    h1 { font-size: 40px; line-height: 1.15; margin: 0 0 12px; }
    h2 { font-size: 28px; margin: 36px 0 12px; }
    p, li { font-size: 20px; }
    .lead { font-size: 22px; background: #E6F4EA; border-left: 8px solid #188038; padding: 14px 16px; }
    a { color: #0b57d0; }
    .links { display: flex; flex-direction: column; gap: 8px; margin: 12px 0 18px; }
    .links a { display: inline-block; background: #fff; border: 1px solid #dadce0; border-radius: 10px; padding: 10px 14px; text-decoration: none; }
    .links a:hover { border-color: #188038; }
    img { width: 100%; height: auto; border: 1px solid #dadce0; border-radius: 12px; margin: 8px 0 4px; background: #fff; }
    .cap { color: #5f6368; font-size: 16px; margin: 0 0 16px; }
    ul, ol { padding-left: 24px; }
    """
    body = ["<main>"]
    for kind, *rest in parts:
        if kind == "h1":
            body.append("<h1>%s</h1>" % html.escape(rest[0]))
        elif kind == "h2":
            body.append("<h2>%s</h2>" % html.escape(rest[0]))
        elif kind == "lead":
            body.append("<p class='lead'>%s</p>" % html.escape(rest[0]))
        elif kind == "p":
            body.append("<p>%s</p>" % html.escape(rest[0]))
        elif kind == "ul":
            body.append("<ul>" + "".join("<li>%s</li>" % html.escape(x) for x in rest[0]) + "</ul>")
        elif kind == "ol":
            body.append("<ol>" + "".join("<li>%s</li>" % html.escape(x) for x in rest[0]) + "</ol>")
        elif kind == "img":
            body.append("<img src='guide/images/%s' alt='%s'><p class='cap'>%s</p>" % (rest[0], html.escape(rest[1]), html.escape(rest[1])))
        elif kind == "links":
            body.append("<div class='links'>" + "".join("<a href='%s'>%s</a>" % (html.escape(url), html.escape(title)) for title, url in rest[0]) + "</div>")
    body.append("</main>")
    page = "<!DOCTYPE html><html lang='ru'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width, initial-scale=1'><title>Как пользоваться пультом РемМонтаж</title><style>%s</style></head><body>%s</body></html>" % (css, "".join(body))
    path = os.path.join(ROOT, "Руководство по пульту.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(page)
    return path


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0B57D0")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "28")
    rpr.append(color)
    rpr.append(underline)
    rpr.append(size)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def render_docx(parts):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(0x20, 0x21, 0x24)
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.6)
        section.right_margin = Cm(1.6)
    for kind, *rest in parts:
        if kind in ("h1", "h2"):
            paragraph = doc.add_heading(rest[0], level=1 if kind == "h1" else 2)
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0x18, 0x80, 0x38) if kind == "h1" else RGBColor(0x20, 0x21, 0x24)
        elif kind in ("p", "lead"):
            doc.add_paragraph(rest[0])
        elif kind == "ul":
            for item in rest[0]:
                doc.add_paragraph(item, style="List Bullet")
        elif kind == "ol":
            for item in rest[0]:
                doc.add_paragraph(item, style="List Number")
        elif kind == "img":
            doc.add_picture(os.path.join(OUT, rest[0]), width=Cm(17.2))
            cap = doc.add_paragraph(rest[1])
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(11)
            cap.runs[0].font.color.rgb = RGBColor(0x5F, 0x63, 0x68)
        elif kind == "links":
            for title, url in rest[0]:
                paragraph = doc.add_paragraph()
                add_hyperlink(paragraph, title, url)
    path = os.path.join(ROOT, "Руководство по пульту.docx")
    doc.save(path)
    return path


if __name__ == "__main__":
    save_diagrams()
    crop_shots()
    parts = blocks()
    print(render_html(parts))
    print(render_docx(parts))
